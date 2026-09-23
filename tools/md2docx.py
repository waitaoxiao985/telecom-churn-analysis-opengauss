# -*- coding: utf-8 -*-
"""md2docx.py -- 把 Markdown 报告转成 Word (.docx)。
支持子集：#/##/###/#### 标题、表格、有序/无序列表、代码块、引用块、
分隔线、**加粗**、`行内代码`、普通段落。中文使用微软雅黑。
用法: python md2docx.py <input.md> <output.docx>
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_run_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), name)


def add_inline(par, text, base_size=10.5, color=None):
    """解析 **加粗** 与 `代码` 行内标记。"""
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = par.add_run(piece[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            set_run_font(run, name="Consolas", size=base_size - 1, color=(165, 39, 39))
        else:
            run = par.add_run(piece)
            set_run_font(run, size=base_size, color=color)


def fill_table_cell(cell, text, bold=False):
    par = cell.paragraphs[0]
    add_inline(par, text, base_size=9)
    for run in par.runs:
        run.font.bold = bold or run.font.bold


def main(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 页面边距
    for sec in doc.sections:
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        # ---- 代码块 ----
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(4)
            run = par.add_run("\n".join(buf))
            set_run_font(run, name="Consolas", size=8.5, color=(40, 40, 40))
            par.paragraph_format.left_indent = Cm(0.5)
            continue

        # ---- 表格 ----
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            rows.pop(1)  # 分隔行
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    if ci < len(table.rows[ri].cells):
                        fill_table_cell(table.rows[ri].cells[ci], cell_text, bold=(ri == 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # ---- 标题 ----
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2))
            if level == 1:
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = par.add_run(text)
                set_run_font(run, size=18, bold=True, color=(31, 78, 121))
            else:
                par = doc.add_paragraph()
                par.paragraph_format.space_before = Pt(10)
                par.paragraph_format.space_after = Pt(4)
                run = par.add_run(text)
                sizes = {2: 15, 3: 12.5, 4: 11}
                set_run_font(run, size=sizes.get(level, 11), bold=True,
                             color=(31, 78, 121) if level == 2 else (47, 84, 150))
            i += 1
            continue

        # ---- 分隔线 ----
        if re.match(r"^-{3,}$", s):
            par = doc.add_paragraph()
            run = par.add_run("─" * 52)
            set_run_font(run, size=8, color=(180, 180, 180))
            i += 1
            continue

        # ---- 引用块 ----
        if s.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            add_inline(par, s.lstrip("> ").strip(), base_size=10,
                       color=(89, 89, 89))
            i += 1
            continue

        # ---- 无序列表 ----
        if re.match(r"^[-*]\s+", s):
            par = doc.add_paragraph(style="List Bullet")
            add_inline(par, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue

        # ---- 有序列表 ----
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            par = doc.add_paragraph(style="List Number")
            add_inline(par, m.group(2))
            i += 1
            continue

        # ---- 空行 ----
        if not s:
            i += 1
            continue

        # ---- 普通段落 ----
        par = doc.add_paragraph()
        add_inline(par, s)

        i += 1

    doc.save(docx_path)
    print("OK ->", docx_path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
